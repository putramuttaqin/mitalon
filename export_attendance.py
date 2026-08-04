from datetime import datetime
import json
import os

from openpyxl import Workbook
from openpyxl.styles import Font


def export_today_attendance(
    attendance_file,
    pegawai_file,
    output_folder="exports"
):
    # Today's key
    today = datetime.now().strftime("%d%m%Y")

    # Load attendance
    with open(attendance_file, "r", encoding="utf-8") as f:
        attendance = json.load(f)

    # Load master pegawai
    with open(pegawai_file, "r", encoding="utf-8") as f:
        pegawai = json.load(f)

    # Attendance today
    hadir = attendance.get(today, {})

    hadir_names = sorted(hadir.keys())
    tidak_hadir = sorted(
        [x for x in pegawai if x not in hadir]
    )

    wb = Workbook()

    ###################################
    # Sheet 1 : Hadir
    ###################################
    ws1 = wb.active
    ws1.title = "Hadir"

    headers = [
        "No",
        "Nama",
        "Jam",
        "Alamat",
        "Koordinat",
        "Synced"
    ]

    for c, h in enumerate(headers, start=1):
        cell = ws1.cell(row=1, column=c)
        cell.value = h
        cell.font = Font(bold=True)

    for i, nama in enumerate(hadir_names, start=2):
        data = hadir[nama]

        timestamp = data.get("timestamp", "")
        jam = ""

        if len(timestamp) >= 15:
            jam = f"{timestamp[9:11]}:{timestamp[11:13]}:{timestamp[13:15]}"

        ws1.cell(i, 1).value = i - 1
        ws1.cell(i, 2).value = nama
        ws1.cell(i, 3).value = jam
        ws1.cell(i, 4).value = data.get("alamat", "")
        ws1.cell(i, 5).value = data.get("koordinat", "")
        ws1.cell(i, 6).value = "Yes" if data.get("synced") else "No"

    ###################################
    # Sheet 2 : Tidak Hadir
    ###################################
    ws2 = wb.create_sheet("Tidak Hadir")

    ws2["A1"] = "No"
    ws2["B1"] = "Nama"

    ws2["A1"].font = Font(bold=True)
    ws2["B1"].font = Font(bold=True)

    for i, nama in enumerate(tidak_hadir, start=2):
        ws2.cell(i, 1).value = i - 1
        ws2.cell(i, 2).value = nama

    ###################################
    # Summary
    ###################################
    ws3 = wb.create_sheet("Summary")

    ws3["A1"] = "Tanggal"
    ws3["B1"] = today

    ws3["A2"] = "Total Pegawai"
    ws3["B2"] = len(pegawai)

    ws3["A3"] = "Hadir"
    ws3["B3"] = len(hadir_names)

    ws3["A4"] = "Tidak Hadir"
    ws3["B4"] = len(tidak_hadir)

    os.makedirs(output_folder, exist_ok=True)

    filename = os.path.join(
        output_folder,
        f"Attendance_{today}.xlsx"
    )

    wb.save(filename)

    return filename

if __name__ == "__main__":
    BASE_DIR = os.path.dirname(os.path.abspath(__file__))

    attendance_file = os.path.join(BASE_DIR, "data", "submissions.json")
    pegawai_file = os.path.join(BASE_DIR, "static", "pegawai.json")
    output_folder = os.path.join(BASE_DIR, "exports")

    filename = export_today_attendance(
        attendance_file=attendance_file,
        pegawai_file=pegawai_file,
        output_folder=output_folder,
    )

    print(f"Attendance report created:")
    print(filename)